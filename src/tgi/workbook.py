"""Build the reviewable workbook: one sheet per functionality, one row per test step.

The JSON export serves tooling. This serves the human who has to sign off a test plan,
so it follows the specification's own numbering and stays sortable and filterable: no
merged cells, a frozen header, an autofilter on every sheet.
"""

from __future__ import annotations

import io
import re
from typing import TYPE_CHECKING, Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from tgi.deliverable import ORPHAN_TESTS, UNNUMBERED, build_deliverable, other_rules_of, tests_by_rule

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet

_HEADER_FILL = PatternFill("solid", start_color="FF0F62FE")
_HEADER_FONT = Font(color="FFFFFFFF", bold=True)
_WRAP = Alignment(vertical="top", wrap_text=True)
_TOP = Alignment(vertical="top")

_TEST_COLUMNS = [
    ("Cas d'utilisation", 18),
    ("Référence règle", 22),
    ("Règle", 60),
    ("Bloc", 10),
    ("Score bloc", 11),
    ("ID test", 12),
    ("Nom du test", 40),
    ("Description", 50),
    ("Étape", 7),
    ("Action", 50),
    ("Résultat attendu", 50),
    ("Statut", 11),
    ("Couvre aussi", 16),
]

_SUMMARY_COLUMNS = [
    ("Fonctionnalité", 16),
    ("Cas d'utilisation", 20),
    ("Règles", 9),
    ("Couvertes", 11),
    ("Couverture", 12),
    ("Tests", 8),
]

_ORPHAN_COLUMNS = [
    ("ID test", 12),
    ("Bloc", 10),
    ("Règles citées", 20),
    ("Nom du test", 45),
    ("Description", 60),
]

# Columns at least this wide hold prose, so their cells wrap
_WRAPPED_COLUMN_WIDTH = 40

# Excel refuses these in a sheet name, and caps it at 31 characters
_FORBIDDEN_IN_SHEET_NAME = re.compile(r"[\[\]:*?/\\]")


def sheet_title(name: str, used: set[str]) -> str:
    """Sheet name Excel accepts, kept unique within the workbook."""
    cleaned = _FORBIDDEN_IN_SHEET_NAME.sub("-", str(name or "Feuille")).strip() or "Feuille"
    candidate = cleaned[:31]
    if candidate not in used:
        used.add(candidate)
        return candidate
    for index in range(2, 100):
        suffix = f" ({index})"
        candidate = cleaned[: 31 - len(suffix)] + suffix
        if candidate not in used:
            used.add(candidate)
            return candidate
    unique = cleaned[:27] + f" {len(used)}"
    used.add(unique)
    return unique


def _write_header(sheet: Worksheet, columns: list[tuple[str, int]]) -> None:
    sheet.append([label for label, _ in columns])
    for index, (_, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _TOP
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"


def _finish_sheet(sheet: Worksheet, columns: list[tuple[str, int]]) -> None:
    """Autofilter over the written range, and wrapped text on the long columns."""
    last_row = max(sheet.max_row, 1)
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last_row}"
    wrapped = {index for index, (_, width) in enumerate(columns, start=1) if width >= _WRAPPED_COLUMN_WIDTH}
    for row in sheet.iter_rows(min_row=2, max_row=last_row):
        for cell in row:
            cell.alignment = _WRAP if cell.column in wrapped else _TOP


def _test_rows(rule: Any, tests: list[dict[str, Any]], use_case: str) -> list[list[Any]]:
    """One row per step, repeating the test columns so filters keep working."""
    rows: list[list[Any]] = []
    for test in tests:
        shared = ", ".join(other_rules_of(test, rule.rule_id))
        steps = [s for s in (test.get("steps") or []) if isinstance(s, dict)] or [{}]
        for step in steps:
            rows.append(
                [
                    use_case,
                    rule.source_ref or rule.rule_id,
                    rule.description,
                    rule.bloc_id,
                    rule.bloc_score,
                    test.get("id", ""),
                    test.get("name", ""),
                    test.get("description", ""),
                    step.get("order", ""),
                    step.get("description", ""),
                    step.get("expected_result", ""),
                    test.get("status", ""),
                    shared,
                ]
            )
    if not tests:
        rows.append(
            [
                use_case,
                rule.source_ref or rule.rule_id,
                rule.description,
                rule.bloc_id,
                rule.bloc_score,
                "",
                "AUCUN TEST",
                "",
                "",
                "",
                "",
                "",
                "",
            ]
        )
    return rows


def build_workbook(rules: list[dict[str, Any]], tests: list[dict[str, Any]], bloc_scores: dict[str, Any]) -> bytes:
    """Return the xlsx bytes for this project."""
    deliverable = build_deliverable(rules, tests, bloc_scores)
    index = tests_by_rule(tests)

    workbook = Workbook()
    used_names: set[str] = set()

    summary = workbook.active
    summary.title = sheet_title("Synthèse", used_names)
    _write_header(summary, _SUMMARY_COLUMNS)

    chapters = list(deliverable["chapters"])
    if deliverable.get("unnumbered"):
        chapters.append(deliverable["unnumbered"])

    for chapter in chapters:
        for group in chapter.groups:
            summary.append(
                [
                    chapter.key,
                    group.title,
                    group.rules_count,
                    group.covered_count,
                    f"{group.coverage_percent} %",
                    group.tests_count,
                ]
            )
    totals = deliverable["totals"]
    summary.append([])
    summary.append(
        ["TOTAL", "", totals["rules"], totals["covered"], f"{totals['coverage_percent']} %", totals["tests"]]
    )
    _finish_sheet(summary, _SUMMARY_COLUMNS)

    for chapter in chapters:
        name = UNNUMBERED if chapter.key == UNNUMBERED else chapter.key
        sheet = workbook.create_sheet(sheet_title(name, used_names))
        _write_header(sheet, _TEST_COLUMNS)
        for group in chapter.groups:
            for rule in group.rules:
                for row in _test_rows(rule, index.get(rule.key, []), group.title):
                    sheet.append(row)
        _finish_sheet(sheet, _TEST_COLUMNS)

    orphans = deliverable["orphan_tests"]
    if orphans:
        sheet = workbook.create_sheet(sheet_title(ORPHAN_TESTS, used_names))
        _write_header(sheet, _ORPHAN_COLUMNS)
        for test in orphans:
            sheet.append(
                [
                    test.get("id", ""),
                    test.get("bloc_id", ""),
                    test.get("business_rule", "") or "aucune",
                    test.get("name", ""),
                    test.get("description", ""),
                ]
            )
        _finish_sheet(sheet, _ORPHAN_COLUMNS)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
