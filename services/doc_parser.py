"""Document parser: Word (.docx), PDF, plain text → clean text."""
from __future__ import annotations

import logging
from pathlib import Path

logger = logging.getLogger(__name__)


class DocParser:
    """Parse various document formats into clean plain text."""

    __slots__ = ()

    def parse(self, path: str | Path) -> str:
        """Auto-detect format and return clean text."""
        p = Path(path)
        suffix = p.suffix.lower()

        if suffix == ".docx":
            return self._parse_docx(p)
        if suffix == ".pdf":
            return self._parse_pdf(p)
        if suffix in {".txt", ".md", ".rst", ""}:
            return self._parse_text(p)

        # Fallback: try text
        logger.warning("Unknown extension %s, trying plain text", suffix)
        return self._parse_text(p)

    def _parse_docx(self, path: Path) -> str:
        from docx import Document  # type: ignore[import-untyped]

        doc = Document(str(path))
        parts: list[str] = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                parts.append(text)

        # Also extract table content
        for table in doc.tables:
            for row in table.rows:
                row_parts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_parts:
                    parts.append(" | ".join(row_parts))

        return "\n\n".join(parts)

    def _parse_pdf(self, path: Path) -> str:
        from pypdf import PdfReader  # type: ignore[import-untyped]

        reader = PdfReader(str(path))
        parts: list[str] = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                parts.append(text.strip())

        return "\n\n".join(parts)

    def _parse_text(self, path: Path) -> str:
        return path.read_text(encoding="utf-8", errors="replace").strip()


doc_parser = DocParser()
