"""Functional tests for the FastAPI app via ASGI transport (no network)."""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path
from typing import TYPE_CHECKING, Any

import pytest
from httpx import ASGITransport, AsyncClient
from opentelemetry import trace

from tgi.config import Settings

if TYPE_CHECKING:
    from collections.abc import AsyncIterator

    from starlette.types import Receive, Scope, Send


@pytest.fixture(autouse=True)
def _reset_tracer() -> None:
    """Reset global tracer provider between tests."""
    trace._TRACER_PROVIDER = None  # type: ignore[attr-defined]
    trace._TRACER_PROVIDER_SET_ONCE._done = False  # type: ignore[attr-defined]


@pytest.fixture
def app_settings(tmp_path: Path) -> Settings:
    return Settings(
        app_name="test_tgi",
        projects_dir=str(tmp_path / "projects"),
        logs=str(tmp_path / "logs"),
        llm_api_key="test-key",
    )


@pytest.fixture
async def client(app_settings: Settings, monkeypatch: pytest.MonkeyPatch) -> AsyncIterator[AsyncClient]:
    # Point the singletons at the temp projects dir
    from tgi.config import settings as global_settings
    from tgi.services import llm as llm_module
    from tgi.services.state_manager import state_manager

    monkeypatch.setattr(global_settings, "projects_dir", app_settings.projects_dir)

    async def _fake_check(self: object, model_id: str, required_tokens: int) -> tuple[bool, int]:
        return True, 0

    async def _fake_chat(self: object, **kwargs: object) -> str:
        return "reponse simulee du QA agent"

    monkeypatch.setattr(llm_module.LLMClient, "check_context_window", _fake_check)
    monkeypatch.setattr(llm_module.LLMClient, "chat", _fake_chat)

    # Sanity: state manager sees the patched dir
    assert state_manager

    from tgi.tgi import create_app

    application = create_app(app_settings=app_settings)
    transport = ASGITransport(app=application)
    async with application.router.lifespan_context(application):
        async with AsyncClient(transport=transport, base_url="http://test") as c:
            yield c


async def _upload_sample(client: AsyncClient) -> str:
    content = b"# Section Un\nUne regle metier importante.\n\n# Section Deux\nUne autre regle."
    files = {"file": ("spec.txt", content, "text/plain")}
    resp = await client.post("/upload", files=files)
    assert resp.status_code == 200
    return resp.json()["project_id"]


async def test_index_empty(client: AsyncClient) -> None:
    resp = await client.get("/")
    assert resp.status_code == 200
    assert "text/html" in resp.headers["content-type"]


async def test_upload_and_project_view(client: AsyncClient) -> None:
    pid = await _upload_sample(client)
    resp = await client.get(f"/projects/{pid}")
    assert resp.status_code == 200

    resp_tests = await client.get(f"/projects/{pid}/tests")
    assert resp_tests.status_code == 200
    assert resp_tests.json() == {"tests": []}


async def test_upload_empty_document_422(client: AsyncClient) -> None:
    files = {"file": ("empty.txt", b"   ", "text/plain")}
    resp = await client.post("/upload", files=files)
    assert resp.status_code == 422


async def test_project_not_found_404(client: AsyncClient) -> None:
    resp = await client.get("/projects/unknown-id")
    assert resp.status_code == 404


async def test_history_endpoint(client: AsyncClient) -> None:
    pid = await _upload_sample(client)
    resp = await client.get(f"/projects/{pid}/history")
    assert resp.status_code == 200
    commits = resp.json()["commits"]
    assert len(commits) >= 1


async def test_validate_split(client: AsyncClient) -> None:
    pid = await _upload_sample(client)
    resp = await client.post(f"/projects/{pid}/validate-split")
    assert resp.status_code == 200
    assert resp.json() == {"status": "ok"}


async def test_chat_uses_mocked_llm(client: AsyncClient) -> None:
    pid = await _upload_sample(client)
    resp = await client.post(f"/projects/{pid}/chat", json={"message": "simple demande"})
    assert resp.status_code == 200
    assert "reponse simulee" in resp.json()["response"]


async def test_chat_empty_message_422(client: AsyncClient) -> None:
    pid = await _upload_sample(client)
    resp = await client.post(f"/projects/{pid}/chat", json={"message": "  "})
    assert resp.status_code == 422


async def test_export_returns_zip(client: AsyncClient) -> None:
    pid = await _upload_sample(client)
    resp = await client.get(f"/projects/{pid}/export")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/zip"
    with zipfile.ZipFile(io.BytesIO(resp.content)) as zf:
        names = zf.namelist()
    assert "state.json" in names
    assert "tests/all_tests.json" in names


async def test_partials_blocs(client: AsyncClient) -> None:
    pid = await _upload_sample(client)
    resp = await client.get(f"/projects/{pid}/partials/blocs")
    assert resp.status_code == 200


async def test_rollback_missing_hash_422(client: AsyncClient) -> None:
    pid = await _upload_sample(client)
    resp = await client.post(f"/projects/{pid}/rollback", json={})
    assert resp.status_code == 422


async def test_run_and_rerun_start(client: AsyncClient, monkeypatch: pytest.MonkeyPatch) -> None:
    from tgi.agents import orchestrator as orch_module

    async def _noop_pipeline(self: object, project_id: str) -> None:
        return None

    async def _noop_rerun(self: object, project_id: str, bloc_id: str) -> None:
        return None

    monkeypatch.setattr(orch_module.Orchestrator, "run_pipeline", _noop_pipeline)
    monkeypatch.setattr(orch_module.Orchestrator, "rerun_bloc", _noop_rerun)

    pid = await _upload_sample(client)
    resp = await client.post(f"/projects/{pid}/run")
    assert resp.status_code == 200
    assert resp.json() == {"status": "started"}

    resp2 = await client.post(f"/projects/{pid}/blocs/bloc-1/rerun")
    assert resp2.status_code == 200
    assert resp2.json()["bloc_id"] == "bloc-1"


async def test_update_test_flow(client: AsyncClient) -> None:
    from tgi.services.state_manager import state_manager

    pid = await _upload_sample(client)
    # Seed a test into the first bloc
    state = await state_manager.load(pid)
    state["blocs"][0]["tests"] = [{"id": "TEST-001", "status": "draft"}]
    await state_manager.save(pid, state)

    resp = await client.put(f"/projects/{pid}/tests/TEST-001", json={"status": "validated"})
    assert resp.status_code == 200
    assert resp.json()["status"] == "validated"

    missing = await client.put(f"/projects/{pid}/tests/NOPE", json={"status": "x"})
    assert missing.status_code == 404


async def test_rollback_success(client: AsyncClient) -> None:
    from tgi.services.git_service import git_service
    from tgi.services.state_manager import state_manager

    pid = await _upload_sample(client)
    first_hash = await git_service.current_hash(pid)
    assert first_hash is not None

    # Make a new commit
    state = await state_manager.load(pid)
    state["validated"] = True
    await state_manager.save(pid, state)
    await git_service.commit(pid, "feat: change")

    resp = await client.post(f"/projects/{pid}/rollback", json={"hash": first_hash})
    assert resp.status_code == 200
    assert resp.json()["status"] == "ok"


async def test_partials_tests_and_history(client: AsyncClient) -> None:
    pid = await _upload_sample(client)
    resp_tests = await client.get(f"/projects/{pid}/partials/tests")
    assert resp_tests.status_code == 200
    resp_hist = await client.get(f"/projects/{pid}/partials/history")
    assert resp_hist.status_code == 200


async def test_blocs_partial_colour_codes_scores(client: AsyncClient) -> None:
    """The blocs partial must show the score and the right colour per outcome."""
    from tgi.services.state_manager import state_manager

    project_id = await _upload_sample(client)
    state = await state_manager.load(project_id)
    state["blocs"] = [
        {
            "id": "b-green",
            "title": "Passe",
            "chunk": "c",
            "rules": [{"id": "R1", "description": "d"}],
            "tests": [],
            "status": "done",
            "score": 97,
            "judge_passes": 1,
            "best_version": 1,
            "judge_history": [{"version": 1, "score": 97, "tests_count": 3}],
        },
        {
            "id": "b-yellow",
            "title": "Revue",
            "chunk": "c",
            "rules": [{"id": "R1", "description": "d"}],
            "tests": [],
            "status": "needs_human",
            "score": 60,
            "judge_passes": 3,
            "best_version": 2,
            "judge_history": [
                {"version": 1, "score": 40, "tests_count": 2},
                {"version": 2, "score": 60, "tests_count": 4},
            ],
        },
        {
            "id": "b-red",
            "title": "Faible",
            "chunk": "c",
            "rules": [{"id": "R1", "description": "d"}],
            "tests": [],
            "status": "needs_human",
            "score": 12,
            "judge_passes": 3,
            "judge_history": [],
        },
        {
            "id": "b-unscored",
            "title": "Non evalue",
            "chunk": "c",
            "rules": [{"id": "R1", "description": "d"}],
            "tests": [],
            "status": "needs_human",
            "score": None,
            "judge_passes": 3,
            "judge_history": [],
        },
        {
            "id": "b-norules",
            "title": "Sans regle",
            "chunk": "c",
            "rules": [],
            "tests": [],
            "status": "done",
            "score": None,
            "judge_passes": 0,
            "judge_history": [],
        },
        {
            "id": "b-error",
            "title": "Casse",
            "chunk": "c",
            "rules": [],
            "tests": [],
            "status": "error",
            "error": "Le modele n'a pas renvoye de JSON exploitable.",
            "judge_passes": 0,
        },
    ]
    await state_manager.save(project_id, state)

    resp = await client.get(f"/projects/{project_id}/partials/blocs")
    assert resp.status_code == 200
    html = resp.text

    # Green: passed the threshold, score shown
    assert "badge-green" in html
    assert "Terminé, score 97%" in html
    # Yellow: kept but flagged, best version reported
    assert "badge-orange" in html
    assert "Revue humaine, score 60%" in html
    assert "version retenue v2" in html
    # Red: score below the bad threshold
    assert "badge-red" in html
    assert "Couverture faible, score 12%" in html
    # Unscored judge and rule-less bloc must not claim a percentage
    assert "score non évalué" in html
    assert "Aucune règle métier" in html
    # Error stays rerunnable
    assert "renvoye de JSON exploitable" in html  # apostrophes are HTML escaped
    assert html.count("↺ Rejouer") == 6
    # The score bar reflects the configured threshold
    assert f"seuil {app_settings_pass_score()}%" in html


def app_settings_pass_score() -> int:
    from tgi.config import Settings

    return Settings().judge_pass_score


async def test_blocs_partial_shows_near_identical_rules(client: AsyncClient) -> None:
    """A reviewer must see the close rules, with both texts, to arbitrate."""
    from tgi.services.state_manager import state_manager

    project_id = await _upload_sample(client)
    state = await state_manager.load(project_id)
    state["blocs"] = [
        {
            "id": "bloc-1",
            "title": "Regles proches",
            "chunk": "c",
            "rules": [
                {"id": "R3", "description": "Si le CDC est Banquier Conseil, supprimer la relation"},
                {"id": "R9", "description": "Si le CDC n'est pas Banquier Conseil, supprimer la relation"},
            ],
            "tests": [],
            "status": "done",
            "score": 100,
            "judge_passes": 1,
            "similar_rules": [{"a": "R3", "b": "R9", "ratio": 0.907}],
        }
    ]
    await state_manager.save(project_id, state)

    resp = await client.get(f"/projects/{project_id}/partials/blocs")
    html = resp.text
    assert "paire(s) de règles très proches" in html
    assert "R3 ~ R9 (91%)" in html
    # Both wordings are shown so the difference is visible
    assert "Si le CDC est Banquier Conseil" in html
    assert "pas Banquier Conseil" in html
    # And the warning is explicit that nothing was merged
    assert "jamais fusionnées" in html


async def _project_with_rules(client: AsyncClient) -> str:
    from tgi.services.state_manager import state_manager

    project_id = await _upload_sample(client)
    state = await state_manager.load(project_id)
    state["blocs"] = [
        {
            "id": "bloc-1",
            "title": "Habilitations",
            "chunk": "c",
            "status": "done",
            "score": 90,
            "judge_passes": 1,
            "rules": [
                {"id": "R1", "source_ref": "VAL01.CU01.RM01", "description": "Le systeme cree une habilitation"},
                {"id": "R2", "source_ref": "", "description": "Sans reference et sans test", "reviewed": True},
            ],
            "tests": [
                {
                    "id": "TEST-001",
                    "bloc_id": "bloc-1",
                    "business_rule": "R1",
                    "name": "creation",
                    "description": "d",
                    "steps": [],
                    "status": "draft",
                    "created_at": "2026-01-01T00:00:00Z",
                    "updated_at": "2026-01-01T00:00:00Z",
                },
                {
                    "id": "TEST-002",
                    "bloc_id": "bloc-1",
                    "business_rule": "R1, R3",
                    "name": "multi",
                    "description": "d",
                    "steps": [],
                    "status": "draft",
                    "created_at": "2026-01-01T00:00:00Z",
                    "updated_at": "2026-01-01T00:00:00Z",
                },
            ],
        },
    ]
    await state_manager.save(project_id, state)
    return project_id


async def test_rules_partial_shows_the_specification_hierarchy(client: AsyncClient) -> None:
    """The deliverable follows the document numbering: functionality, use case, rule."""
    project_id = await _project_with_rules(client)
    resp = await client.get(f"/projects/{project_id}/partials/rules")
    assert resp.status_code == 200
    html = resp.text

    assert "VAL01" in html  # functionality level
    assert "VAL01.CU01" in html  # use case level
    assert "RM01" in html  # rule label taken from the reference
    assert "Hors numérotation" in html  # the rule without a reference has its chapter
    assert "2 tests" in html  # R1 is covered by TEST-001 and by the multi rule TEST-002
    assert "badge-orange" in html  # R2 is not
    # Tests are loaded on expansion, not inlined
    assert "/rules/R1/tests" in html
    assert 'hx-trigger="toggle once"' in html
    # TEST-002 cites R1 and R3: R1 exists, so it is shared coverage, not an orphan
    assert "Tests non rattachés" not in html


async def test_rules_partial_surfaces_orphan_tests(client: AsyncClient) -> None:
    """A test citing only rules that do not exist in its bloc must be visible."""
    from tgi.services.state_manager import state_manager

    project_id = await _project_with_rules(client)
    state = await state_manager.load(project_id)
    state["blocs"][0]["tests"].append(
        {
            "id": "TEST-099",
            "bloc_id": "bloc-1",
            "business_rule": "R42",
            "name": "cite une regle inexistante",
            "description": "d",
            "steps": [],
            "status": "draft",
            "created_at": "2026-01-01T00:00:00Z",
            "updated_at": "2026-01-01T00:00:00Z",
        }
    )
    await state_manager.save(project_id, state)

    html = (await client.get(f"/projects/{project_id}/partials/rules")).text
    assert "Tests non rattachés" in html
    assert "TEST-099" in html
    assert "R42" in html


async def test_rules_partial_filters_server_side(client: AsyncClient) -> None:
    project_id = await _project_with_rules(client)
    only_uncovered = (await client.get(f"/projects/{project_id}/partials/rules?uncovered=1")).text
    assert "Sans reference et sans test" in only_uncovered
    assert "Le systeme cree une habilitation" not in only_uncovered

    # The haystack includes the bloc title on purpose, so filter on a rule specific word
    by_text = (await client.get(f"/projects/{project_id}/partials/rules?q=cree")).text
    assert "Le systeme cree une habilitation" in by_text
    assert "Sans reference et sans test" not in by_text

    by_reference = (await client.get(f"/projects/{project_id}/partials/rules?q=VAL01.CU01")).text
    assert "Le systeme cree une habilitation" in by_reference

    nothing = (await client.get(f"/projects/{project_id}/partials/rules?q=zzzintrouvable")).text
    assert "Aucune règle ne correspond au filtre" in nothing


async def test_rule_tests_fragment_lists_the_covering_tests(client: AsyncClient) -> None:
    project_id = await _project_with_rules(client)
    html = (await client.get(f"/projects/{project_id}/blocs/bloc-1/rules/R1/tests")).text
    assert "TEST-001" in html
    assert "TEST-002" in html
    # TEST-002 also covers R3, shown so the reader does not count it twice
    assert "couvre aussi" in html
    assert "R3" in html

    empty = (await client.get(f"/projects/{project_id}/blocs/bloc-1/rules/R2/tests")).text
    assert "Aucun test pour cette règle" in empty


async def test_rule_can_be_edited_and_marked_reviewed(client: AsyncClient) -> None:
    project_id = await _project_with_rules(client)
    resp = await client.put(
        f"/projects/{project_id}/blocs/bloc-1/rules/R1",
        json={"description": "Formulation corrigee", "source_ref": "VAL01.CU01.RM09", "reviewed": True},
    )
    assert resp.status_code == 200
    assert resp.json()["description"] == "Formulation corrigee"

    rules = (await client.get(f"/projects/{project_id}/rules")).json()["rules"]
    edited = next(r for r in rules if r["id"] == "R1")
    assert edited["source_ref"] == "VAL01.CU01.RM09"
    assert edited["reviewed"] is True
    assert edited["bloc_id"] == "bloc-1"


async def test_rule_edit_is_scoped_to_its_bloc(client: AsyncClient) -> None:
    """R1 exists in every bloc: the pair (bloc, rule) is the identity."""
    project_id = await _project_with_rules(client)
    missing = await client.put(f"/projects/{project_id}/blocs/bloc-99/rules/R1", json={"reviewed": True})
    assert missing.status_code == 404
    unknown = await client.put(f"/projects/{project_id}/blocs/bloc-1/rules/R404", json={"reviewed": True})
    assert unknown.status_code == 404


async def test_tests_partial_filters_and_paginates_server_side(client: AsyncClient) -> None:
    """A test card is about 5 kB of HTML: filtering and paging must happen server side."""
    project_id = await _project_with_rules(client)
    html = (await client.get(f"/projects/{project_id}/partials/tests")).text
    assert "pour 2 règles" in html
    assert "TEST-001" in html and "TEST-002" in html
    assert 'value="R1"' in html and 'value="R3"' in html
    # The JSON payload lives in a single quoted attribute, otherwise it closes it early
    assert "x-data='testEditor(" in html
    assert 'x-data="testEditor(' not in html

    by_rule = (await client.get(f"/projects/{project_id}/partials/tests?rule=R3")).text
    assert "TEST-002" in by_rule
    assert "TEST-001" not in by_rule
    assert "1 correspondent au filtre" in by_rule

    paged = (await client.get(f"/projects/{project_id}/partials/tests?per_page=1")).text
    assert "1 à 1 sur 2" in paged
    assert "page 1/2" in paged

    second = (await client.get(f"/projects/{project_id}/partials/tests?per_page=1&page=2")).text
    assert "TEST-002" in second
    assert "TEST-001" not in second

    out_of_range = (await client.get(f"/projects/{project_id}/partials/tests?per_page=1&page=99")).text
    assert "page 2/2" in out_of_range  # clamped, never an empty page

    none = (await client.get(f"/projects/{project_id}/partials/tests?q=zzzintrouvable")).text
    assert "Aucun test ne correspond au filtre" in none


async def test_progress_partial_reports_the_run(client: AsyncClient) -> None:
    from tgi.services.state_manager import state_manager

    project_id = await _project_with_rules(client)
    state = await state_manager.load(project_id)
    state["blocs"] = state["blocs"] + [
        {"id": "bloc-2", "title": "b", "chunk": "c", "status": "pending", "rules": [], "tests": []},
        {"id": "bloc-3", "title": "c", "chunk": "c", "status": "error", "rules": [], "tests": []},
    ]
    await state_manager.save(project_id, state)

    html = (await client.get(f"/projects/{project_id}/partials/progress")).text
    assert "2/3</strong> blocs (67 %)" in html
    assert "erreur 1" in html
    assert "en attente 1" in html
    assert "2 règles" in html and "2 tests" in html


async def test_export_contains_the_reviewable_workbook(client: AsyncClient) -> None:
    import io as _io
    import zipfile as _zipfile

    from openpyxl import load_workbook

    project_id = await _project_with_rules(client)
    resp = await client.get(f"/projects/{project_id}/export")
    assert resp.status_code == 200

    with _zipfile.ZipFile(_io.BytesIO(resp.content)) as archive:
        names = archive.namelist()
        assert "tests.xlsx" in names
        assert "tests/all_tests.json" in names  # JSON kept for tooling
        workbook = load_workbook(_io.BytesIO(archive.read("tests.xlsx")))

    assert workbook.sheetnames[0].startswith("Synth")
    assert "VAL01" in workbook.sheetnames
    assert "Hors numérotation" in workbook.sheetnames
    sheet = workbook["VAL01"]
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref is not None
    assert [c.value for c in sheet[1]][:3] == ["Cas d'utilisation", "Référence règle", "Règle"]


async def test_fragments_are_compressed(client: AsyncClient) -> None:
    """The rules tree is 1.53 MB for 850 rules and gzips 22 times smaller."""
    project_id = await _project_with_rules(client)

    plain = await client.get(f"/projects/{project_id}/partials/rules", headers={"Accept-Encoding": "identity"})
    zipped = await client.get(f"/projects/{project_id}/partials/rules", headers={"Accept-Encoding": "gzip"})
    assert plain.status_code == zipped.status_code == 200
    assert zipped.headers.get("content-encoding") == "gzip"
    # httpx decodes transparently, so compare the wire length the server reported
    assert int(zipped.headers["content-length"]) < len(plain.content)
    assert "rules-panel" in zipped.text  # and it still decodes to the same page


async def test_the_event_stream_is_never_compressed() -> None:
    """gzip buffers a stream, so live progress would arrive in bursts.

    Driven at the ASGI layer: opening a real SSE connection would never close.
    """
    from tgi.tgi import ConditionalGZipMiddleware

    seen: list[str] = []

    async def app(scope: Scope, receive: Receive, send: Send) -> None:
        seen.append(scope["path"])
        await send({"type": "http.response.start", "status": 200, "headers": [(b"content-type", b"text/plain")]})
        await send({"type": "http.response.body", "body": b"x" * 5000})

    async def receive() -> dict[str, Any]:
        return {"type": "http.request", "body": b"", "more_body": False}

    async def run(path: str) -> list[tuple[bytes, bytes]]:
        headers: list[tuple[bytes, bytes]] = []

        async def send(message: dict[str, Any]) -> None:
            if message["type"] == "http.response.start":
                headers.extend(message["headers"])

        scope: Scope = {
            "type": "http",
            "path": path,
            "method": "GET",
            "headers": [(b"accept-encoding", b"gzip")],
        }
        await ConditionalGZipMiddleware(app)(scope, receive, send)
        return headers

    stream_headers = await run("/projects/abc/stream")
    assert not any(name == b"content-encoding" for name, _ in stream_headers)

    fragment_headers = await run("/projects/abc/partials/rules")
    assert (b"content-encoding", b"gzip") in fragment_headers


async def test_filter_dropdowns_are_ordered_numerically(client: AsyncClient) -> None:
    """bloc-10 used to sit between bloc-1 and bloc-2 in the filter lists."""
    from tgi.services.state_manager import state_manager

    project_id = await _project_with_rules(client)
    state = await state_manager.load(project_id)
    template = state["blocs"][0]
    for bloc_id in ("bloc-2", "bloc-9", "bloc-10", "bloc-20"):
        state["blocs"].append(
            {
                **template,
                "id": bloc_id,
                "rules": [{"id": "R9", "source_ref": "", "description": f"regle de {bloc_id}"}],
                "tests": [
                    {
                        "id": f"TEST-{bloc_id}",
                        "bloc_id": bloc_id,
                        "business_rule": "R9",
                        "name": "t",
                        "description": "d",
                        "steps": [],
                        "status": "draft",
                        "created_at": "2026-01-01T00:00:00Z",
                        "updated_at": "2026-01-01T00:00:00Z",
                    }
                ],
            }
        )
    await state_manager.save(project_id, state)

    html = (await client.get(f"/projects/{project_id}/partials/tests")).text
    blocs = re.findall(r'name="bloc"(.*?)</select>', html, re.S)[0]
    assert re.findall(r'value="(bloc-[\d]+)"', blocs) == ["bloc-1", "bloc-2", "bloc-9", "bloc-10", "bloc-20"]

    tree = (await client.get(f"/projects/{project_id}/partials/rules")).text
    # Group titles of the unnumbered chapter, in the order the tree renders them
    groups = re.findall(r'font-mono text-sm text-gray-200">(bloc-\d+) ·', tree)
    assert groups == ["bloc-1", "bloc-2", "bloc-9", "bloc-10", "bloc-20"]
