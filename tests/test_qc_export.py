"""Tests for the QC export: one sheet, columns and TYPE derivation (FR-NEW-034 to 038)."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from tgi.qc_export import build_qc_workbook

_HEADER = [
    "Subject",
    "Test Name",
    "Description",
    "Classification",
    "Step Name",
    "Step Description",
    "Expected Results",
]


def _state(refs: list[str], **extra: Any) -> dict[str, Any]:
    state: dict[str, Any] = {
        "containers": {"F03.EU05.CU01": "Déléguer temporairement"},
        "requirements": [
            {"ref": "F03.EU05.CU01.RM01"},
            {"ref": "F03.EU05.CU01.RM02"},
            {"ref": "F03.EU05.CU01.EM01"},
            {"ref": "F03.EU05.CU01.EM02"},
            {"ref": "E04.M01"},
            {"ref": "E04.M02"},
            {"ref": "E04.N01"},
            {"ref": "E04.N02"},
            {"ref": "E04.T01"},
            {"ref": "E04.T02"},
        ],
        "scenarios": [
            {
                "id": "SC-001",
                "container": "F03.EU05.CU01",
                "tests": [
                    {
                        "id": "TEST-0001",
                        "name": "cas nominal",
                        "description": "Une description.",
                        "requirement_refs": refs,
                        "steps": [
                            {"order": 1, "description": "agir", "expected_result": "vu"},
                            {"order": 2, "description": "vérifier", "expected_result": "correct"},
                        ],
                    }
                ],
            }
        ],
    }
    state.update(extra)
    return state


def test_header_row_is_exact() -> None:
    content, _ = build_qc_workbook(_state(["F03.EU05.CU01.RM01"]))
    sheet = load_workbook(BytesIO(content))["QC"]
    assert [cell.value for cell in sheet[1]] == _HEADER


def test_one_row_per_step_with_subject_and_test_name_repeated() -> None:
    content, _ = build_qc_workbook(_state(["F03.EU05.CU01.RM01"]))
    sheet = load_workbook(BytesIO(content))["QC"]
    assert sheet.max_row == 3  # header + 2 steps
    assert sheet["A2"].value == sheet["A3"].value
    assert sheet["B2"].value == sheet["B3"].value
    assert sheet["E2"].value == "Étape 1"
    assert sheet["E3"].value == "Étape 2"


def test_subject_and_test_name_composition() -> None:
    content, _ = build_qc_workbook(_state(["F03.EU05.CU01.RM01"]))
    sheet = load_workbook(BytesIO(content))["QC"]
    assert sheet["A2"].value == "F03.EU05.CU01-RM"
    assert sheet["B2"].value == "TRA_TEST-0001_cas nominal__F03.EU05.CU01.RM01"
    assert "Exigences validées : F03.EU05.CU01.RM01" in sheet["C2"].value


def test_test_name_composition_is_locked_regression() -> None:
    """FR-NEW-068 / DEC-021: Test Name stays TRA_<id>_<name>__<refs>, not rewritten for the email example."""
    content, _ = build_qc_workbook(_state(["F03.EU05.CU01.RM01", "F03.EU05.CU01.RM02"]))
    sheet = load_workbook(BytesIO(content))["QC"]
    assert sheet["B2"].value == "TRA_TEST-0001_cas nominal__F03.EU05.CU01.RM01, F03.EU05.CU01.RM02"


def test_description_composition_is_locked_regression() -> None:
    """FR-NEW-069: Description stays <description>\\nExigences validées : <refs>."""
    content, _ = build_qc_workbook(_state(["F03.EU05.CU01.RM01"]))
    sheet = load_workbook(BytesIO(content))["QC"]
    assert sheet["C2"].value == "Une description.\nExigences validées : F03.EU05.CU01.RM01"


def test_subject_follows_onglet_type_name_for_ihm_and_rm() -> None:
    content, _ = build_qc_workbook(_state(["E04.M01"]))
    sheet = load_workbook(BytesIO(content))["QC"]
    assert sheet["A2"].value == "IHM_E04"

    content, _ = build_qc_workbook(_state(["F03.EU05.CU01.RM01"]))
    sheet = load_workbook(BytesIO(content))["QC"]
    assert sheet["A2"].value == "F03.EU05.CU01-RM"


def test_classification_column_present_after_description() -> None:
    content, _ = build_qc_workbook(_state(["E04.M01"]))
    sheet = load_workbook(BytesIO(content))["QC"]
    headers = [cell.value for cell in sheet[1]]
    assert headers.index("Classification") == headers.index("Description") + 1
    assert sheet["D2"].value == "MOA"


def test_subject_derivation_uses_onglet_type_name_for_each_kind() -> None:
    cases = {
        "F03.EU05.CU01.RM01": "F03.EU05.CU01-RM",
        "F03.EU05.CU01.EM01": "F03.EU05.CU01-EMOE",
        "E04.M01": "IHM_E04",
        "E04.N01": "IHM_E04",
    }
    for ref, expected_subject in cases.items():
        content, warnings = build_qc_workbook(_state([ref]))
        sheet = load_workbook(BytesIO(content))["QC"]
        assert sheet["A2"].value == expected_subject, ref
        assert warnings == []


def test_a_cell_beyond_32767_characters_is_truncated() -> None:
    long_description = "x" * 40000
    state = _state(["F03.EU05.CU01.RM01"])
    state["scenarios"][0]["tests"][0]["description"] = long_description

    content, _ = build_qc_workbook(state)
    sheet = load_workbook(BytesIO(content))["QC"]
    assert len(sheet["C2"].value) == 32767
    assert sheet["C2"].value.endswith("…")


def test_a_test_with_no_requirement_ref_still_produces_a_row_subject_inconnu() -> None:
    content, warnings = build_qc_workbook(_state([]))
    sheet = load_workbook(BytesIO(content))["QC"]
    assert sheet["A2"].value == "INCONNU"
    assert any("sans référence" in w or "inconnu" in w.lower() for w in warnings)
