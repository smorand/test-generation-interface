"""Tests for the document parser."""

from __future__ import annotations

from pathlib import Path

from tgi.services.doc_parser import DocParser


def test_parse_txt(tmp_path: Path) -> None:
    f = tmp_path / "spec.txt"
    f.write_text("  Regle metier en texte.  ", encoding="utf-8")
    assert DocParser().parse(f) == "Regle metier en texte."


def test_parse_md(tmp_path: Path) -> None:
    f = tmp_path / "spec.md"
    f.write_text("# Titre\ncontenu", encoding="utf-8")
    assert "Titre" in DocParser().parse(f)


def test_parse_unknown_extension_falls_back_to_text(tmp_path: Path) -> None:
    f = tmp_path / "spec.xyz"
    f.write_text("contenu brut", encoding="utf-8")
    assert DocParser().parse(f) == "contenu brut"


def test_parse_no_extension(tmp_path: Path) -> None:
    f = tmp_path / "README"
    f.write_text("sans extension", encoding="utf-8")
    assert DocParser().parse(f) == "sans extension"


def test_parse_accepts_string_path(tmp_path: Path) -> None:
    f = tmp_path / "spec.txt"
    f.write_text("string path", encoding="utf-8")
    assert DocParser().parse(str(f)) == "string path"


def test_parse_docx(tmp_path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Premier paragraphe.")
    doc.add_paragraph("Deuxieme paragraphe.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "cellA"
    table.rows[0].cells[1].text = "cellB"
    path = tmp_path / "spec.docx"
    doc.save(str(path))

    result = DocParser().parse(path)
    assert "Premier paragraphe." in result
    assert "Deuxieme paragraphe." in result
    assert "cellA | cellB" in result


def _docx_with_structure(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("Section principale", level=2)
    doc.add_paragraph("Une regle metier importante.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "colonne A"
    table.rows[0].cells[1].text = "colonne B"
    doc.add_heading("Deuxieme section", level=3)
    doc.add_paragraph("Une autre regle.")
    doc.save(str(path))


def test_docx_preserves_headings_as_markdown(tmp_path: Path) -> None:
    path = tmp_path / "structure.docx"
    _docx_with_structure(path)
    text = DocParser().parse(path)
    assert "## Section principale" in text
    assert "### Deuxieme section" in text


def test_docx_keeps_tables_in_document_order(tmp_path: Path) -> None:
    """A table must stay under its own heading, not be appended at the end."""
    path = tmp_path / "order.docx"
    _docx_with_structure(path)
    text = DocParser().parse(path)
    table_pos = text.index("colonne A | colonne B")
    second_heading_pos = text.index("### Deuxieme section")
    assert table_pos < second_heading_pos


def test_docx_skips_hidden_comment_paragraphs(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "comments.docx"
    doc = Document()
    doc.add_paragraph("Contenu de specification.")
    style = doc.styles.add_style("commentaire cache", 1)
    doc.add_paragraph("Remarque interne a ignorer", style=style)
    doc.save(str(path))

    text = DocParser().parse(path)
    assert "Contenu de specification." in text
    assert "Remarque interne" not in text
